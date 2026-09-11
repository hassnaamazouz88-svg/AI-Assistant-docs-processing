from pathlib import Path
from typing import Any

from docx import Document


def generate_docx(
    content_plan: dict[str, Any],
    output_path: str | Path
) -> str:
    """
    Génère un fichier DOCX à partir d'une structure
    de contenu produite par ContentPlanner.

    Parameters
    ----------
    content_plan:
        Dictionnaire contenant le titre et les sections
        à générer.

    output_path:
        Chemin du fichier DOCX à créer.

    Returns
    -------
    str:
        Chemin du fichier DOCX généré.
    """

    # Vérification du contenu
    if not isinstance(content_plan, dict):
        raise ValueError(
            "Le contenu à générer doit être un dictionnaire."
        )

    if "titre" not in content_plan:
        raise ValueError(
            "Le contenu ne contient pas de champ 'titre'."
        )

    if "sections" not in content_plan:
        raise ValueError(
            "Le contenu ne contient pas de champ 'sections'."
        )

    titre = content_plan["titre"]
    sections = content_plan["sections"]

    if not isinstance(titre, str):
        raise ValueError(
            "Le champ 'titre' doit être une chaîne de caractères."
        )

    if not isinstance(sections, list):
        raise ValueError(
            "Le champ 'sections' doit être une liste."
        )

    # Préparer le chemin de sortie
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # Créer le document Word
    document = Document()

    # Titre principal
    document.add_heading(
        titre,
        level=0
    )

    # Ajouter les sections
    for section in sections:

        if not isinstance(section, dict):
            raise ValueError(
                "Chaque section doit être un dictionnaire."
            )

        section_title = section.get("titre")
        section_content = section.get("contenu")

        if section_title is None:
            raise ValueError(
                "Une section ne contient pas de titre."
            )

        if section_content is None:
            raise ValueError(
                f"La section '{section_title}' "
                "ne contient pas de contenu."
            )

        # Titre de la section
        document.add_heading(
            section_title,
            level=1
        )

        # Contenu de la section
        document.add_paragraph(
            section_content
        )

    # Sauvegarder le document
    document.save(output_path)

    return str(output_path)