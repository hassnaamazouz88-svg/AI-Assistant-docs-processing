from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def generate_xlsx(
    content_plan: dict[str, Any],
    output_path: str | Path
) -> str:
    """
    Génère un fichier XLSX à partir d'une structure
    de contenu produite par ContentPlanner.

    Parameters
    ----------
    content_plan:
        Dictionnaire contenant le titre et les sections
        à générer.

    output_path:
        Chemin du fichier XLSX à créer.

    Returns
    -------
    str:
        Chemin du fichier XLSX généré.
    """

    # Vérification du contenu
    if not isinstance(content_plan, dict):
        raise ValueError(
            "Le contenu à générer doit être un dictionnaire."
        )

    if "titre" not in content_plan:
        raise ValueError(
            "Le contenu ne contient pas de champ 'titre'."
        )

    if "sections" not in content_plan:
        raise ValueError(
            "Le contenu ne contient pas de champ 'sections'."
        )

    titre = content_plan["titre"]
    sections = content_plan["sections"]

    if not isinstance(titre, str):
        raise ValueError(
            "Le champ 'titre' doit être une chaîne de caractères."
        )

    if not isinstance(sections, list):
        raise ValueError(
            "Le champ 'sections' doit être une liste."
        )

    # Préparer le chemin de sortie
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # Créer le classeur Excel
    workbook = Workbook()

    # Feuille principale
    worksheet = workbook.active
    worksheet.title = "Synthèse"

    # Titre principal
    worksheet["A1"] = titre
    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    # En-têtes du tableau
    worksheet["A3"] = "Section"
    worksheet["B3"] = "Contenu"

    worksheet["A3"].font = Font(bold=True)
    worksheet["B3"].font = Font(bold=True)

    entete_fill = PatternFill(start_color="E8E8EC", end_color="E8E8EC", fill_type="solid")
    worksheet["A3"].fill = entete_fill
    worksheet["B3"].fill = entete_fill

    # Ajouter les sections
    row = 4

    for section in sections:

        if not isinstance(section, dict):
            raise ValueError(
                "Chaque section doit être un dictionnaire."
            )

        section_title = section.get("titre")
        section_content = section.get("contenu")

        if section_title is None:
            raise ValueError(
                "Une section ne contient pas de titre."
            )

        if section_content is None:
            raise ValueError(
                f"La section '{section_title}' "
                "ne contient pas de contenu."
            )

        worksheet.cell(row=row,column=1,value=section_title)

        worksheet.cell(row=row,column=2,value=section_content)

        row += 1

    # Mise en forme
    bordure = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0")
    )

    for current_row in worksheet.iter_rows(
        min_row=3,
        max_row=row - 1,
        min_col=1,
        max_col=2
    ):
        for cell in current_row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )
            # Sans bordure, l'aperçu PDF (converti via LibreOffice) ne
            # ressemble à rien d'autre qu'à du texte brut — les traits
            # de cellule sont ce qui donne visuellement l'apparence
            # d'un tableau une fois converti.
            cell.border = bordure

    # Largeur des colonnes
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 70

    # Mise en page pour l'impression / la conversion PDF : sans ces
    # réglages, la colonne B (large) ne tient pas en largeur de page et
    # LibreOffice imprime "Section" sur une page puis "Contenu" sur la
    # suivante, ce qui donne l'impression que ce n'est pas un tableau.
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_area = f"A1:B{row - 1}"

    # Figer les en-têtes
    worksheet.freeze_panes = "A4"

    # Sauvegarder le fichier
    workbook.save(output_path)

    return str(output_path)